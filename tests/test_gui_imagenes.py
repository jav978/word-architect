"""Pruebas de la inserción de imágenes desde la interfaz gráfica."""

import struct
import tempfile
import unittest
import zlib
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document

from src.gui import (
    _bloque_imagen,
    _convertir_a_bloques,
    _es_linea_imagen,
    _partes_imagen,
    insertar_imagen,
)
from src.documento import crear_documento_profesional


def _generar_png(ruta: Path, ancho: int = 10, alto: int = 10) -> None:
    """Generar un archivo PNG válido y mínimo para las pruebas."""
    firma = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", ancho, alto, 8, 2, 0, 0, 0)
    ihdr_crc = zlib.crc32(b"IHDR" + ihdr) & 0xFFFFFFFF
    ihdr_chunk = struct.pack(">I", 13) + b"IHDR" + ihdr + struct.pack(">I", ihdr_crc)
    fila = b"\x00" + b"\x00\x00\x00" * ancho
    comprimido = zlib.compress(fila * alto)
    idat_crc = zlib.crc32(b"IDAT" + comprimido) & 0xFFFFFFFF
    idat_chunk = struct.pack(">I", len(comprimido)) + b"IDAT" + comprimido + struct.pack(">I", idat_crc)
    iend_crc = zlib.crc32(b"IEND") & 0xFFFFFFFF
    iend_chunk = struct.pack(">I", 0) + b"IEND" + struct.pack(">I", iend_crc)
    ruta.write_bytes(firma + ihdr_chunk + idat_chunk + iend_chunk)


class TestConversionImagenes(unittest.TestCase):
    """Verificar que el editor reconoce y valida las referencias de imagen."""

    def test_es_linea_imagen(self):
        """Detectar correctamente las líneas de referencia de imagen."""
        self.assertTrue(_es_linea_imagen("{{imagen: /tmp/foto.png | 4 | Foto}}"))
        self.assertTrue(_es_linea_imagen("{{imagen: /tmp/foto.png}}"))
        self.assertFalse(_es_linea_imagen("Imagen: /tmp/foto.png"))
        self.assertFalse(_es_linea_imagen("{{imagen: /tmp/foto.png"))  # sin cierre
        self.assertFalse(_es_linea_imagen("Párrafo normal con {{imagen: x}}"))  # no empieza por el marcador

    def test_partes_imagen_solo_ruta(self):
        """Extraer solo la ruta cuando no hay ancho ni leyenda."""
        ruta, ancho, leyenda = _partes_imagen("{{imagen: /home/usuario/foto.png}}")
        self.assertEqual(ruta, "/home/usuario/foto.png")
        self.assertEqual(ancho, "5.5")
        self.assertEqual(leyenda, "")

    def test_partes_imagen_con_ancho_y_leyenda(self):
        """Extraer ruta, ancho y leyenda."""
        ruta, ancho, leyenda = _partes_imagen("{{imagen: /tmp/foto.png | 3.5 | Foto de familia}}")
        self.assertEqual(ruta, "/tmp/foto.png")
        self.assertEqual(ancho, "3.5")
        self.assertEqual(leyenda, "Foto de familia")

    def test_partes_imagen_ruta_con_espacios_y_acentos(self):
        """La ruta conserva espacios, acentos y caracteres especiales."""
        ruta, ancho, leyenda = _partes_imagen("{{imagen: /tmp/foto de vacación 2026.png | 4 | Verano}}")
        self.assertEqual(ruta, "/tmp/foto de vacación 2026.png")
        self.assertEqual(ancho, "4")
        self.assertEqual(leyenda, "Verano")

    def test_convierte_referencia_imagen_en_bloque(self):
        """Una línea de imagen produce un único bloque image."""
        with tempfile.TemporaryDirectory() as tmp:
            ruta_imagen = Path(tmp) / "foto.png"
            _generar_png(ruta_imagen)
            bloques = _convertir_a_bloques(f"{{{{imagen: {ruta_imagen} | 4 | Foto de prueba}}}}")

        self.assertEqual(len(bloques), 1)
        self.assertEqual(bloques[0]["tipo"], "image")
        self.assertEqual(bloques[0]["ruta"], str(ruta_imagen))
        self.assertEqual(bloques[0]["ancho"], 4.0)
        self.assertEqual(bloques[0]["leyenda"], "Foto de prueba")

    def test_convierte_imagen_sin_leyenda(self):
        """Una imagen sin leyenda genera el bloque sin el campo 'leyenda'."""
        with tempfile.TemporaryDirectory() as tmp:
            ruta_imagen = Path(tmp) / "foto.png"
            _generar_png(ruta_imagen)
            bloques = _convertir_a_bloques(f"{{{{imagen: {ruta_imagen} | 3}}}}")

        self.assertEqual(bloques[0]["tipo"], "image")
        self.assertEqual(bloques[0]["ancho"], 3.0)
        self.assertNotIn("leyenda", bloques[0])

    def test_orden_texto_tabla_imagen(self):
        """El orden entre párrafos, tablas e imágenes se conserva."""
        with tempfile.TemporaryDirectory() as tmp:
            ruta_imagen = Path(tmp) / "foto.png"
            _generar_png(ruta_imagen)
            contenido = "\n".join(
                [
                    "Introducción",
                    "| Columna | Valor |",
                    "| A | 1 |",
                    f"{{{{imagen: {ruta_imagen} | 4 | Foto}}}}",
                    "Cierre",
                ]
            )
            bloques = _convertir_a_bloques(contenido)

        tipos = [b["tipo"] for b in bloques]
        self.assertEqual(tipos, ["paragraph", "table", "image", "paragraph"])

    def test_rechaza_imagen_inexistente(self):
        """Una ruta que no existe lanza ValueError."""
        with self.assertRaises(ValueError) as contexto:
            _bloque_imagen("{{imagen: /tmp/no_existe.png | 4 | Foto}}")
        self.assertIn("No se encontró la imagen", str(contexto.exception))

    def test_rechaza_formato_incompatible(self):
        """Un archivo con extensión no reconocida lanza ValueError."""
        with tempfile.TemporaryDirectory() as tmp:
            ruta_texto = Path(tmp) / "documento.txt"
            ruta_texto.write_text("esto no es una imagen")
            with self.assertRaises(ValueError) as contexto:
                _bloque_imagen(f"{{{{imagen: {ruta_texto} | 4 | Foto}}}}")
        self.assertIn("no compatible", str(contexto.exception))

    def test_rechaza_ancho_invalido(self):
        """Un ancho no numérico o menor o igual a cero lanza ValueError."""
        with tempfile.TemporaryDirectory() as tmp:
            ruta_imagen = Path(tmp) / "foto.png"
            _generar_png(ruta_imagen)
            with self.assertRaises(ValueError):
                _bloque_imagen(f"{{{{imagen: {ruta_imagen} | abc | Foto}}}}")
            with self.assertRaises(ValueError):
                _bloque_imagen(f"{{{{imagen: {ruta_imagen} | 0 | Foto}}}}")
            with self.assertRaises(ValueError):
                _bloque_imagen(f"{{{{imagen: {ruta_imagen} | -2 | Foto}}}}")

    def test_rechaza_ruta_vacia(self):
        """Una referencia sin ruta lanza ValueError."""
        with self.assertRaises(ValueError):
            _bloque_imagen("{{imagen: | 4 | Foto}}")


class TestDocumentoConImagen(unittest.TestCase):
    """Generar archivos DOCX reales que contengan la imagen."""

    def setUp(self):
        self.temporal = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporal.cleanup)
        self.directorio = Path(self.temporal.name)

    def _crear_png(self, nombre: str) -> Path:
        ruta = self.directorio / nombre
        _generar_png(ruta)
        return ruta

    def test_genera_docx_con_imagen(self):
        """El bloque image incrusta la imagen físicamente en el DOCX."""
        ruta_imagen = self._crear_png("foto.png")
        bloques = _convertir_a_bloques(f"{{{{imagen: {ruta_imagen} | 2 | Foto de prueba}}}}")

        ruta = crear_documento_profesional(
            "Informe con imagen",
            bloques,
            self.directorio / "informe.docx",
        )

        self.assertTrue(ruta.exists())
        documento = Document(ruta)
        self.assertEqual(len(documento.inline_shapes), 1)

    def test_ancho_y_leyenda_en_docx(self):
        """La leyenda y el ancho configurados se incorporan al documento."""
        ruta_imagen = self._crear_png("foto.png")
        bloques = _convertir_a_bloques(f"{{{{imagen: {ruta_imagen} | 3 | Leyenda de la imagen}}}}")

        ruta = crear_documento_profesional(
            "Informe",
            bloques,
            self.directorio / "informe.docx",
        )

        documento = Document(ruta)
        self.assertEqual(len(documento.inline_shapes), 1)
        # 3 párrafos: título (Heading 1), imagen y leyenda (Caption)
        self.assertEqual(len(documento.paragraphs), 3)
        self.assertEqual(documento.paragraphs[2].text, "Leyenda de la imagen")
        self.assertEqual(documento.paragraphs[2].style.name, "Caption")

    def test_imagen_entre_texto_y_tabla(self):
        """La imagen mantiene su posición respecto a otros bloques en el DOCX."""
        ruta_imagen = self._crear_png("foto.png")
        contenido = "\n".join(
            [
                "Primera sección",
                f"{{{{imagen: {ruta_imagen} | 2 | Foto del medio}}}}",
                "| Columna | Valor |",
                "| A | 1 |",
                "Última sección",
            ]
        )
        bloques = _convertir_a_bloques(contenido)
        tipos = [b["tipo"] for b in bloques]
        self.assertEqual(tipos, ["paragraph", "image", "table", "paragraph"])

        ruta = crear_documento_profesional(
            "Orden",
            bloques,
            self.directorio / "orden.docx",
        )
        documento = Document(ruta)
        self.assertEqual(len(documento.inline_shapes), 1)
        self.assertEqual(len(documento.tables), 1)

    def test_ruta_inexistente_no_crea_archivo(self):
        """Si la imagen no existe, no se genera el DOCX."""
        ruta_destino = self.directorio / "invalido.docx"
        bloques = [{"tipo": "image", "ruta": "/tmp/no_existe.png", "ancho": 4}]

        with self.assertRaises(FileNotFoundError):
            crear_documento_profesional("Inválido", bloques, ruta_destino)

        self.assertEqual(list(self.directorio.iterdir()), [])

    def test_anchoy_no_numero_no_crea_bloque(self):
        """Un ancho no numérico impide la creación del bloque."""
        with tempfile.TemporaryDirectory() as tmp:
            ruta_imagen = Path(tmp) / "foto.png"
            _generar_png(ruta_imagen)
            with self.assertRaises(ValueError):
                _convertir_a_bloques(f"{{{{imagen: {ruta_imagen} | ancho | Foto}}}}")


class TestInsercionImagenGUI(unittest.TestCase):
    """Verificar que insertar_imagen no modifica el contenido al cancelar."""

    def test_cancelar_selector_no_modifica_contenido(self):
        """Si el usuario cancela el selector de archivos, el contenido no cambia."""
        mock_cuadro = MagicMock()
        mock_cuadro.index.return_value = "1.0"

        with patch("src.gui.cuadro_contenido", mock_cuadro), \
             patch("src.gui.filedialog.askopenfilename", return_value=""):
            insertar_imagen()

        mock_cuadro.insert.assert_not_called()


if __name__ == "__main__":
    unittest.main()
