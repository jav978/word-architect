"""Pruebas del generador profesional de Word-Architect."""

import tempfile
import unittest
from pathlib import Path

from docx import Document

from src.documento import (
    ConfiguracionDocumento,
    crear_documento_profesional,
)


class TestGeneradorProfesional(unittest.TestCase):
    """Verificar las funciones avanzadas del generador."""

    def setUp(self):
        self.temporal = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporal.cleanup)
        self.directorio = Path(self.temporal.name)

    def test_crea_documento_con_tabla_y_titulos(self):
        """Comprobar títulos, párrafos y tablas editables."""

        ruta = crear_documento_profesional(
            "Informe de prueba",
            [
                {
                    "tipo": "heading",
                    "nivel": 1,
                    "texto": "Introducción",
                },
                {
                    "tipo": "paragraph",
                    "texto": "Este es un párrafo de prueba.",
                },
                {
                    "tipo": "table",
                    "encabezados": ["Materia", "Nota"],
                    "filas": [
                        ["Matemática", "18"],
                        ["Programación", "20"],
                    ],
                },
            ],
            self.directorio / "informe.docx",
        )

        documento = Document(ruta)

        self.assertEqual(
            documento.paragraphs[0].text,
            "Informe de prueba",
        )
        self.assertIn(
            ("Heading 1", "Introducción"),
            [
                (p.style.name, p.text)
                for p in documento.paragraphs
            ],
        )

        self.assertEqual(len(documento.tables), 1)
        self.assertEqual(
            documento.tables[0].cell(1, 0).text,
            "Matemática",
        )
        self.assertEqual(
            documento.tables[0].cell(2, 1).text,
            "20",
        )

    def test_configuracion_profesional(self):
        """Comprobar fuente, tamaño e interlineado."""

        configuracion = ConfiguracionDocumento(
            tamano_pagina="A4",
            orientacion="vertical",
            fuente="Arial",
            tamano_fuente=12,
            interlineado=1.5,
        )

        ruta = crear_documento_profesional(
            "Trabajo académico",
            [{"tipo": "paragraph", "texto": "Contenido de prueba."}],
            self.directorio / "academico.docx",
            configuracion=configuracion,
        )

        documento = Document(ruta)
        seccion = documento.sections[0]

        self.assertLess(
            seccion.page_width,
            seccion.page_height,
        )
        self.assertEqual(
            documento.styles["Normal"].font.name,
            "Arial",
        )
        self.assertEqual(
            documento.styles["Normal"].font.size.pt,
            12,
        )
        self.assertEqual(
            documento.styles["Normal"].paragraph_format.line_spacing,
            1.5,
        )

    def test_secciones_vertical_y_horizontal(self):
        """Comprobar que el documento puede combinar orientaciones."""

        ruta = crear_documento_profesional(
            "Documento con dos orientaciones",
            [
                {
                    "tipo": "paragraph",
                    "texto": "Primera sección vertical.",
                },
                {
                    "tipo": "section_break",
                    "orientacion": "horizontal",
                },
                {
                    "tipo": "paragraph",
                    "texto": "Segunda sección horizontal.",
                },
            ],
            self.directorio / "secciones.docx",
        )

        documento = Document(ruta)

        self.assertEqual(len(documento.sections), 2)

        primera = documento.sections[0]
        segunda = documento.sections[1]

        self.assertLess(
            primera.page_width,
            primera.page_height,
        )
        self.assertGreater(
            segunda.page_width,
            segunda.page_height,
        )

    def test_rechaza_filas_de_tabla_inconsistentes(self):
        """No permitir filas con un número incorrecto de celdas."""

        with self.assertRaises(ValueError):
            crear_documento_profesional(
                "Tabla inválida",
                [
                    {
                        "tipo": "table",
                        "encabezados": ["A", "B"],
                        "filas": [["Valor único"]],
                    }
                ],
                self.directorio / "invalido.docx",
            )

    def test_no_sobrescribe_documentos(self):
        """Conservar el archivo anterior al repetir el nombre."""

        destino = self.directorio / "informe.docx"

        primera = crear_documento_profesional(
            "Informe",
            [{"tipo": "paragraph", "texto": "Versión inicial."}],
            destino,
        )

        segunda = crear_documento_profesional(
            "Informe",
            [{"tipo": "paragraph", "texto": "Segunda versión."}],
            destino,
        )

        self.assertEqual(primera.name, "informe.docx")
        self.assertEqual(segunda.name, "informe_2.docx")

        self.assertIn(
            "Versión inicial.",
            [p.text for p in Document(primera).paragraphs],
        )
        self.assertIn(
            "Segunda versión.",
            [p.text for p in Document(segunda).paragraphs],
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)