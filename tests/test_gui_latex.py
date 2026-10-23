"""Pruebas de soporte LaTeX → OMML (ecuaciones nativas de Word).

Verifica:
  * Fórmulas en línea $...$
  * Ecuaciones independientes $$...$$
  * Conservación del texto circundante
  * Generación de DOCX real con estructuras OMML nativas
  * Que las ecuaciones NO se convierten en imágenes
  * Manejo de expresiones inválidas
  * Compatibilidad con las 57 pruebas anteriores
"""

import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from src.documento import crear_documento_profesional
from src.gui import _convertir_a_bloques
from src.ecuaciones import latex_a_omml, latex_a_omml_parrafo, LatexConversionError

OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _tiene_omml(ruta: Path) -> bool:
    """Detectar si el DOCX contiene elementos OMML (ecuaciones nativas)."""
    with zipfile.ZipFile(ruta) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    return "oMath" in xml


def _tiene_imagen(ruta: Path) -> bool:
    with zipfile.ZipFile(ruta) as z:
        nombres = z.namelist()
    return any(n.startswith("word/media/") for n in nombres)


class TestConversionLatexOmml(unittest.TestCase):
    """Conversión de expresiones LaTeX a elementos OMML."""

    def test_en_linea_simple(self):
        omml = latex_a_omml("f(x)=x^2+3x+2")
        self.assertIsNotNone(omml)
        self.assertEqual(omml.tag, f"{{{OMML_NS}}}oMath")

    def test_fraccion(self):
        omml = latex_a_omml(r"\frac{a+b}{c}")
        self.assertEqual(omml.tag, f"{{{OMML_NS}}}f")

    def test_potencia(self):
        omml = latex_a_omml("x^2")
        self.assertEqual(omml.tag, f"{{{OMML_NS}}}sSup")

    def test_subindice(self):
        omml = latex_a_omml("a_i")
        self.assertEqual(omml.tag, f"{{{OMML_NS}}}sSub")

    def test_raiz(self):
        omml = latex_a_omml(r"\sqrt{x}")
        self.assertEqual(omml.tag, f"{{{OMML_NS}}}rad")

    def test_integral(self):
        omml = latex_a_omml(r"\int_0^1 x^2\,dx")
        self.assertEqual(omml.tag, f"{{{OMML_NS}}}oMath")

    def test_sumatoria(self):
        omml = latex_a_omml(r"\sum_{i=1}^{n} i")
        self.assertEqual(omml.tag, f"{{{OMML_NS}}}oMath")

    def test_limite(self):
        omml = latex_a_omml(r"\lim_{x \to \infty} f(x)")
        self.assertEqual(omml.tag, f"{{{OMML_NS}}}oMath")

    def test_omml_parrafo_display(self):
        omml_para = latex_a_omml_parrafo(r"\int_0^1 x^2\,dx")
        self.assertEqual(omml_para.tag, f"{{{OMML_NS}}}oMathPara")

    def test_expresion_invalida_lanza_error(self):
        with self.assertRaises(LatexConversionError):
            latex_a_omml(r"\comando_inexistente{x}")


class TestBloquesDesdeEditor(unittest.TestCase):
    """Conversión del contenido del editor a bloques con ecuaciones."""

    def test_ecuacion_independiente(self):
        bloques = _convertir_a_bloques("$$\\int_0^1 x^2\\,dx=\\frac{1}{3}$$")
        self.assertEqual(len(bloques), 1)
        self.assertEqual(bloques[0]["tipo"], "equation")
        self.assertEqual(bloques[0]["modo"], "display")
        self.assertEqual(
            bloques[0]["latex"], r"\int_0^1 x^2\,dx=\frac{1}{3}"
        )

    def test_ecuacion_en_linea(self):
        bloques = _convertir_a_bloques("La función es $f(x)=x^2+3x+2$.")
        self.assertEqual(len(bloques), 1)
        self.assertEqual(bloques[0]["tipo"], "paragraph")
        partes = bloques[0]["partes"]
        self.assertEqual(len(partes), 3)
        self.assertEqual(partes[0], {"tipo": "texto", "contenido": "La función es "})
        self.assertEqual(
            partes[1],
            {"tipo": "ecuacion", "modo": "inline", "latex": "f(x)=x^2+3x+2"},
        )
        self.assertEqual(partes[2], {"tipo": "texto", "contenido": "."})

    def test_varias_ecuaciones_en_linea(self):
        bloques = _convertir_a_bloques("$a$ y $b$ son variables.")
        partes = bloques[0]["partes"]
        self.assertEqual(len(partes), 4)
        self.assertEqual(partes[0], {"tipo": "ecuacion", "modo": "inline", "latex": "a"})
        self.assertEqual(partes[1], {"tipo": "texto", "contenido": " y "})
        self.assertEqual(partes[2], {"tipo": "ecuacion", "modo": "inline", "latex": "b"})
        self.assertEqual(partes[3], {"tipo": "texto", "contenido": " son variables."})

    def test_texto_tabla_ecuacion_conserva_orden(self):
        bloques = _convertir_a_bloques(
            "Introducción\n"
            "$$\sum_{i=1}^{n} i = \\frac{n(n+1)}{2}$$\n"
            "| Columna | Valor |\n"
            "| A | 1 |\n"
            "Conclusión"
        )
        tipos = [b["tipo"] for b in bloques]
        self.assertEqual(tipos, ["paragraph", "equation", "table", "paragraph"])

    def test_barra_sola_no_es_ecuacion(self):
        """Una línea con una sola $ sin cierre no se trata como ecuación."""
        bloques = _convertir_a_bloques("Precio $ 100")
        self.assertEqual(bloques[0]["tipo"], "paragraph")
        self.assertEqual(bloques[0]["texto"], "Precio $ 100")

    def test_ecuacion_vacia_display_no_valida(self):
        """Una línea $$...$$ vacía no genera bloque equation válido."""
        bloques = _convertir_a_bloques("$$$$")
        self.assertEqual(len(bloques), 1)
        self.assertEqual(bloques[0]["tipo"], "paragraph")
        self.assertEqual(bloques[0]["texto"], "$$$$")


class TestDocumentoConEcuaciones(unittest.TestCase):
    """Generación de DOCX reales con ecuaciones OMML."""

    def setUp(self):
        self.temporal = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporal.cleanup)
        self.directorio = Path(self.temporal.name)

    def test_genera_docx_con_ecuacion_display(self):
        """El bloque equation produce OMML nativo en el DOCX."""
        ruta = crear_documento_profesional(
            "Ecuación",
            _convertir_a_bloques("$$\\int_0^1 x^2\\,dx=\\frac{1}{3}$$"),
            self.directorio / "ecuacion.docx",
        )
        self.assertTrue(ruta.exists())
        self.assertTrue(_tiene_omml(ruta))
        self.assertFalse(_tiene_imagen(ruta))

        documento = Document(ruta)
        self.assertGreaterEqual(len(documento.paragraphs), 2)  # título + ecuación

    def test_genera_docx_con_ecuacion_en_linea(self):
        """La ecuación en línea se incrusta dentro del párrafo."""
        ruta = crear_documento_profesional(
            "Inline",
            _convertir_a_bloques("La función es $f(x)=x^2+3x+2$."),
            self.directorio / "inline.docx",
        )
        self.assertTrue(_tiene_omml(ruta))
        self.assertFalse(_tiene_imagen(ruta))

        documento = Document(ruta)
        # Al menos un párrafo debe contener OMML en su XML
        omml_en_parrafos = any(
            OMML_NS in p._p.xml for p in documento.paragraphs
        )
        self.assertTrue(omml_en_parrafos)

    def test_conserva_texto_alrededor_ecuacion(self):
        """El texto antes y después de la ecuación en línea se conserva."""
        ruta = crear_documento_profesional(
            "Contexto",
            _convertir_a_bloques("Inicio $x^2$ fin."),
            self.directorio / "contexto.docx",
        )
        documento = Document(ruta)
        texto_completo = "\n".join(p.text for p in documento.paragraphs)
        self.assertIn("Inicio", texto_completo)
        self.assertIn("fin", texto_completo)

    def test_ecuacion_entre_texto_y_tabla(self):
        """El orden texto → ecuación → tabla se preserva en el DOCX."""
        ruta = crear_documento_profesional(
            "Orden",
            _convertir_a_bloques(
                "Primera sección\n"
                "$$\sum_{i=1}^{n} i$$\n"
                "| Col | Val |\n"
                "| A | 1 |\n"
                "Última sección"
            ),
            self.directorio / "orden.docx",
        )
        documento = Document(ruta)
        self.assertEqual(len(documento.tables), 1)
        self.assertTrue(_tiene_omml(ruta))
        self.assertFalse(_tiene_imagen(ruta))

    def test_expresion_invalida_no_genera_archivo(self):
        """Una expresión LaTeX inválida impide la generación del DOCX."""
        ruta = self.directorio / "invalido.docx"
        # \comando_inexistente no es un comando soportado
        bloques = [{"tipo": "equation", "modo": "display", "latex": r"\comando_inexistente{x}"}]
        with self.assertRaises(ValueError):
            crear_documento_profesional("Inválido", bloques, ruta)
        self.assertEqual(list(self.directorio.iterdir()), [])

    def test_ecuacion_sin_latex_no_genera_archivo(self):
        """Un bloque equation sin campo 'latex' lanza ValueError."""
        ruta = self.directorio / "vacio.docx"
        bloques = [{"tipo": "equation", "modo": "display"}]
        with self.assertRaises(ValueError):
            crear_documento_profesional("Vacío", bloques, ruta)
        self.assertEqual(list(self.directorio.iterdir()), [])


class TestCompatibilidadAnterior(unittest.TestCase):
    """Verificar que las funcionalidades previas (tablas, imágenes, texto) siguen funcionando."""

    def setUp(self):
        self.temporal = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporal.cleanup)
        self.directorio = Path(self.temporal.name)

    def test_tabla_sigue_funcionando(self):
        ruta = crear_documento_profesional(
            "Tabla",
            _convertir_a_bloques("| A | B |\n| 1 | 2 |"),
            self.directorio / "tabla.docx",
        )
        documento = Document(ruta)
        self.assertEqual(len(documento.tables), 1)
        self.assertFalse(_tiene_omml(ruta))

    def test_texto_plano_sigue_funcionando(self):
        ruta = crear_documento_profesional(
            "Texto",
            _convertir_a_bloques("Párrafo simple sin fórmulas."),
            self.directorio / "texto.docx",
        )
        documento = Document(ruta)
        self.assertEqual(documento.paragraphs[1].text, "Párrafo simple sin fórmulas.")
        self.assertFalse(_tiene_omml(ruta))


if __name__ == "__main__":
    unittest.main()
