"""
Pruebas automatizadas del generador Word-Architect.

Comprueba que los documentos se crean correctamente,
que el texto se conserva y que los estilos funcionan.
"""

import tempfile
import unittest
from pathlib import Path

from docx import Document
from src.gui import crear_documento


class TestGeneradorDocumento(unittest.TestCase):
    """Pruebas básicas del generador de documentos Word."""

    def setUp(self):
        """Preparar una carpeta temporal para cada prueba."""

        self.directorio_temporal = tempfile.TemporaryDirectory()
        self.carpeta = Path(self.directorio_temporal.name)

    def tearDown(self):
        """Eliminar únicamente los archivos temporales de las pruebas."""

        self.directorio_temporal.cleanup()

    def test_crear_documento_con_texto(self):
        """Comprobar que el título y los párrafos se conservan."""

        destino = self.carpeta / "prueba_texto.docx"

        titulo = "Prueba de Word-Architect"
        parrafos = [
            "Este es el primer párrafo.",
            "La informática permite procesar información.",
            "Este documento es editable.",
        ]

        ruta = crear_documento(
            titulo,
            parrafos,
            destino,
        )

        # Comprobar que el archivo existe.
        self.assertTrue(ruta.exists())
        self.assertEqual(ruta.suffix, ".docx")

        # Abrir el documento generado.
        documento = Document(str(ruta))

        textos = [
            parrafo.text
            for parrafo in documento.paragraphs
        ]

        # Comprobar que no se perdió información.
        self.assertIn(titulo, textos)

        for texto in parrafos:
            self.assertIn(texto, textos)

    def test_titulos_y_listas_editables(self):
        """Comprobar la conversión de marcas a estilos de Word."""

        destino = self.carpeta / "prueba_formatos.docx"

        contenido = [
            "# Introducción",
            "Texto de introducción.",
            "## Desarrollo",
            "- Primer elemento",
            "- Segundo elemento",
            "1. Primer paso",
        ]

        ruta = crear_documento(
            "Prueba de formatos",
            contenido,
            destino,
        )

        documento = Document(str(ruta))

        estilos_por_texto = {
            parrafo.text: parrafo.style.name
            for parrafo in documento.paragraphs
        }

        # Verificar que los títulos sean estilos de Word.
        self.assertEqual(
            estilos_por_texto["Introducción"],
            "Heading 1",
        )

        self.assertEqual(
            estilos_por_texto["Desarrollo"],
            "Heading 2",
        )

        # Verificar las listas.
        self.assertEqual(
            estilos_por_texto["Primer elemento"],
            "List Bullet",
        )

        self.assertEqual(
            estilos_por_texto["Primer paso"],
            "List Number",
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)