"""Generación y guardado de documentos Word para Word-Architect.

Este módulo concentra la lógica de formato y guardado. Los módulos ``main.py``
y ``gui.py`` pueden seguir utilizando ``crear_documento`` para la generación
básica. Para documentos más completos se ofrece ``crear_documento_profesional``
con soporte opcional para tablas, imágenes, encabezados, logotipos, numeración
de páginas y secciones con orientación vertical u horizontal.

La API básica mantiene su comportamiento anterior para no romper las pruebas ni
los documentos creados desde versiones previas del proyecto.
"""

from __future__ import annotations

import re
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional, Union

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

if __package__:
    from .ecuaciones import (  # type: ignore[no-redef]
        latex_a_omml,
        latex_a_omml_parrafo,
        LatexConversionError,
        M_NS as _ECUACIONES_M_NS,
    )
else:
    from ecuaciones import (  # type: ignore[no-redef]
        latex_a_omml,
        latex_a_omml_parrafo,
        LatexConversionError,
        M_NS as _ECUACIONES_M_NS,
    )


# ==========================================
# CONFIGURACIÓN DEL PROYECTO
# ==========================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = PROJECT_ROOT / "output"

RUTA_DESTINO = Union[str, Path]
PATRON_NUMERADO = re.compile(r"^\d+[.)]\s+")

MENSAJE_TITULO_VACIO = "El título no puede estar vacío."
MENSAJE_SIN_PARRAFOS = "Debes escribir al menos un párrafo."
MENSAJE_PARRAFOS_INVALIDOS = "Los párrafos deben ser una colección de textos."
MENSAJE_PARRAFO_INVALIDO = "Cada párrafo debe ser un texto."
MENSAJE_TITULO_INVALIDO = "El título debe ser un texto."


@dataclass
class ConfiguracionDocumento:
    """Opciones de formato para ``crear_documento_profesional``.

    Las medidas de los márgenes y del logotipo se expresan en pulgadas.
    ``orientacion`` acepta ``"vertical"``/``"portrait"`` o
    ``"horizontal"``/``"landscape"``. ``tamano_pagina`` acepta A4, Letter
    (Carta) o Legal.
    """

    tamano_pagina: str = "A4"
    orientacion: str = "vertical"
    fuente: str = "Arial"
    tamano_fuente: float = 11
    interlineado: float = 1.15
    margen_superior: float = 0.8
    margen_inferior: float = 0.8
    margen_izquierdo: float = 1.0
    margen_derecho: float = 1.0
    espacio_despues_parrafo: float = 8
    texto_encabezado: Optional[str] = None
    texto_pie: Optional[str] = "Generado con Word-Architect"
    numeracion_paginas: bool = True
    ruta_logotipo: Optional[RUTA_DESTINO] = None
    ancho_logotipo: float = 0.65
    autor: Optional[str] = None
    asunto: Optional[str] = None
    palabras_clave: Optional[str] = None
    crear_portada: bool = False
    subtitulo_portada: Optional[str] = None
    institucion: Optional[str] = None


# ==========================================
# VALIDACIÓN
# ==========================================

def _validar_entrada(
    titulo: str,
    parrafos: Iterable[str],
) -> tuple[str, list[str]]:
    """Comprobar título y párrafos y devolverlos normalizados."""

    if not isinstance(titulo, str):
        raise TypeError(MENSAJE_TITULO_INVALIDO)

    titulo = titulo.strip()

    if not titulo:
        raise ValueError(MENSAJE_TITULO_VACIO)

    if isinstance(parrafos, str) or not isinstance(parrafos, Iterable):
        raise TypeError(MENSAJE_PARRAFOS_INVALIDOS)

    textos: list[str] = []

    for parrafo in parrafos:
        if not isinstance(parrafo, str):
            raise TypeError(MENSAJE_PARRAFO_INVALIDO)
        textos.append(parrafo)

    if not any(texto.strip() for texto in textos):
        raise ValueError(MENSAJE_SIN_PARRAFOS)

    return titulo, textos


def _validar_bloques(titulo: str, bloques: Iterable[Any]) -> tuple[str, list[Any]]:
    """Validar el título y los bloques estructurados del generador profesional."""

    if not isinstance(titulo, str):
        raise TypeError(MENSAJE_TITULO_INVALIDO)

    titulo = titulo.strip()
    if not titulo:
        raise ValueError(MENSAJE_TITULO_VACIO)

    if isinstance(bloques, (str, bytes)) or not isinstance(bloques, Iterable):
        raise TypeError("Los bloques deben ser una colección de elementos.")

    elementos = list(bloques)
    if not elementos:
        raise ValueError(MENSAJE_SIN_PARRAFOS)

    for bloque in elementos:
        if isinstance(bloque, str):
            continue
        if not isinstance(bloque, Mapping):
            raise TypeError(
                "Cada bloque debe ser un texto o un diccionario con su tipo y contenido."
            )

    return titulo, elementos


def _validar_configuracion(configuracion: ConfiguracionDocumento) -> None:
    """Validar opciones que podrían causar errores al crear el documento."""

    if not isinstance(configuracion, ConfiguracionDocumento):
        raise TypeError("La configuración debe ser una instancia de ConfiguracionDocumento.")

    tamano = configuracion.tamano_pagina.strip().lower()
    if tamano not in {"a4", "letter", "carta", "legal"}:
        raise ValueError("El tamaño de página debe ser A4, Letter/Carta o Legal.")

    orientacion = configuracion.orientacion.strip().lower()
    if orientacion not in {"vertical", "portrait", "horizontal", "landscape"}:
        raise ValueError("La orientación debe ser vertical o horizontal.")

    if not configuracion.fuente.strip():
        raise ValueError("La fuente no puede estar vacía.")

    if configuracion.tamano_fuente <= 0:
        raise ValueError("El tamaño de fuente debe ser mayor que cero.")

    if configuracion.interlineado <= 0:
        raise ValueError("El interlineado debe ser mayor que cero.")

    margenes = (
        configuracion.margen_superior,
        configuracion.margen_inferior,
        configuracion.margen_izquierdo,
        configuracion.margen_derecho,
    )
    if any(margen < 0 for margen in margenes):
        raise ValueError("Los márgenes no pueden ser negativos.")

    if configuracion.ancho_logotipo <= 0:
        raise ValueError("El ancho del logotipo debe ser mayor que cero.")

    if configuracion.ruta_logotipo is not None:
        ruta_logo = Path(configuracion.ruta_logotipo)
        if not ruta_logo.is_file():
            raise FileNotFoundError(f"No se encontró el logotipo: {ruta_logo}")


# ==========================================
# NOMBRES Y RUTAS
# ==========================================

def _nombre_archivo(titulo: str) -> str:
    """Convertir el título en un nombre de archivo seguro."""

    nombre = re.sub(r"[^\w\s-]", "", titulo.lower())
    nombre = re.sub(r"[\s-]+", "_", nombre).strip("_")

    return nombre or "documento"


def _ruta_libre(ruta: Path) -> Path:
    """Devolver una ruta que no exista, añadiendo ``_2``, ``_3``…"""

    if not ruta.exists():
        return ruta

    contador = 2

    while True:
        candidata = ruta.with_name(f"{ruta.stem}_{contador}{ruta.suffix}")
        if not candidata.exists():
            return candidata
        contador += 1


def _resolver_destino(
    titulo: str,
    ruta_destino: Optional[RUTA_DESTINO],
) -> Path:
    """Calcular la ruta final sin sobrescribir archivos existentes."""

    if ruta_destino is None:
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        return _ruta_libre(OUTPUT_DIR / f"{_nombre_archivo(titulo)}.docx")

    ruta = Path(ruta_destino)

    if ruta.exists() and ruta.is_dir():
        raise IsADirectoryError(f"La ruta de destino es una carpeta: {ruta}")

    # Conservar extensiones distintas: ``informe.v2`` -> ``informe.v2.docx``.
    if ruta.suffix.lower() != ".docx":
        ruta = ruta.with_name(f"{ruta.name}.docx")

    if ruta.exists() and ruta.is_dir():
        raise IsADirectoryError(f"La ruta de destino es una carpeta: {ruta}")

    ruta.parent.mkdir(parents=True, exist_ok=True)

    return _ruta_libre(ruta)


# ==========================================
# FORMATO BASE (API EXISTENTE)
# ==========================================

def _aplicar_formato(documento: Document):
    """Aplicar el formato base anterior y devolver la primera sección."""

    seccion = documento.sections[0]
    seccion.top_margin = Inches(0.8)
    seccion.bottom_margin = Inches(0.8)
    seccion.left_margin = Inches(1)
    seccion.right_margin = Inches(1)

    estilo_normal = documento.styles["Normal"]
    estilo_normal.font.name = "Arial"
    estilo_normal.font.size = Pt(11)

    for nombre_estilo in ("Heading 1", "Heading 2", "Heading 3"):
        documento.styles[nombre_estilo].font.name = "Arial"

    documento.styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)

    return seccion


def _agregar_titulo(documento: Document, titulo: str) -> None:
    """Agregar el título principal con el formato de cabecera."""

    encabezado = documento.add_heading(titulo, level=1)
    encabezado.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if not encabezado.runs:
        return

    run = encabezado.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(31, 78, 121)


def _agregar_parrafo(documento: Document, texto: str) -> None:
    """Agregar un párrafo normal con separación inferior."""

    parrafo = documento.add_paragraph(texto)
    parrafo.paragraph_format.space_after = Pt(8)


def _agregar_contenido(
    documento: Document,
    parrafos: Iterable[str],
    procesar_marcas: bool,
) -> None:
    """Agregar párrafos, interpretando marcas solo si se solicitan."""

    for texto in parrafos:
        texto = texto.strip()

        if not texto:
            continue

        if not procesar_marcas:
            _agregar_parrafo(documento, texto)
            continue

        if texto in {"#", "##", "###", "-", "*"}:
            continue

        if texto.startswith("### "):
            documento.add_heading(texto[4:].strip(), level=3)
        elif texto.startswith("## "):
            documento.add_heading(texto[3:].strip(), level=2)
        elif texto.startswith("# "):
            documento.add_heading(texto[2:].strip(), level=1)
        elif texto.startswith(("- ", "* ")):
            documento.add_paragraph(texto[2:].strip(), style="List Bullet")
        elif PATRON_NUMERADO.match(texto):
            elemento = PATRON_NUMERADO.sub("", texto, count=1)
            documento.add_paragraph(elemento, style="List Number")
        else:
            _agregar_parrafo(documento, texto)


# ==========================================
# FUNCIONES PROFESIONALES: PÁGINA Y ESTILOS
# ==========================================

def _dimensiones_pagina(tamano_pagina: str) -> tuple[float, float]:
    """Devolver ancho y alto en pulgadas para un tamaño admitido."""

    tamano = tamano_pagina.strip().lower()
    dimensiones = {
        "a4": (8.27, 11.69),
        "letter": (8.5, 11.0),
        "carta": (8.5, 11.0),
        "legal": (8.5, 14.0),
    }
    return dimensiones[tamano]


def _configurar_seccion(
    seccion,
    configuracion: ConfiguracionDocumento,
    orientacion: Optional[str] = None,
) -> None:
    """Aplicar tamaño, orientación y márgenes a una sección de Word."""

    ancho, alto = _dimensiones_pagina(configuracion.tamano_pagina)
    orientacion_final = (orientacion or configuracion.orientacion).strip().lower()

    if orientacion_final in {"horizontal", "landscape"}:
        seccion.orientation = WD_ORIENT.LANDSCAPE
        seccion.page_width = Inches(alto)
        seccion.page_height = Inches(ancho)
    else:
        seccion.orientation = WD_ORIENT.PORTRAIT
        seccion.page_width = Inches(ancho)
        seccion.page_height = Inches(alto)

    seccion.top_margin = Inches(configuracion.margen_superior)
    seccion.bottom_margin = Inches(configuracion.margen_inferior)
    seccion.left_margin = Inches(configuracion.margen_izquierdo)
    seccion.right_margin = Inches(configuracion.margen_derecho)


def _aplicar_estilos_profesionales(
    documento: Document,
    configuracion: ConfiguracionDocumento,
) -> None:
    """Configurar fuente, tamaños y espaciado de estilos de Word."""

    estilo_normal = documento.styles["Normal"]
    estilo_normal.font.name = configuracion.fuente
    estilo_normal.font.size = Pt(configuracion.tamano_fuente)
    estilo_normal.paragraph_format.line_spacing = configuracion.interlineado
    estilo_normal.paragraph_format.space_after = Pt(
        configuracion.espacio_despues_parrafo
    )

    for nombre_estilo in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        documento.styles[nombre_estilo].font.name = configuracion.fuente

    documento.styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
    documento.styles["Heading 2"].font.color.rgb = RGBColor(46, 87, 126)
    documento.styles["Heading 3"].font.color.rgb = RGBColor(70, 70, 70)


def _agregar_campo_word(parrafo, instruccion: str) -> None:
    """Insertar un campo dinámico de Word, por ejemplo PAGE o NUMPAGES."""

    run = parrafo.add_run()

    comienzo = OxmlElement("w:fldChar")
    comienzo.set(qn("w:fldCharType"), "begin")

    texto_instruccion = OxmlElement("w:instrText")
    texto_instruccion.set(qn("xml:space"), "preserve")
    texto_instruccion.text = f" {instruccion} "

    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")

    texto_inicial = OxmlElement("w:t")
    texto_inicial.text = "1"

    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")

    run._r.append(comienzo)
    run._r.append(texto_instruccion)
    run._r.append(separador)
    run._r.append(texto_inicial)
    run._r.append(fin)


def _configurar_encabezado_pie(
    documento: Document,
    configuracion: ConfiguracionDocumento,
) -> None:
    """Agregar encabezado, logotipo, pie y numeración dinámica."""

    seccion = documento.sections[0]

    if configuracion.ruta_logotipo is not None or configuracion.texto_encabezado:
        encabezado = seccion.header
        parrafo = encabezado.paragraphs[0]
        parrafo.text = ""

        if configuracion.ruta_logotipo is not None:
            parrafo.add_run().add_picture(
                str(configuracion.ruta_logotipo),
                width=Inches(configuracion.ancho_logotipo),
            )
            if configuracion.texto_encabezado:
                parrafo.add_run("   ")

        if configuracion.texto_encabezado:
            run = parrafo.add_run(configuracion.texto_encabezado)
            run.font.name = configuracion.fuente
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(90, 90, 90)

    pie = seccion.footer
    parrafo_pie = pie.paragraphs[0]
    parrafo_pie.text = ""
    parrafo_pie.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if configuracion.texto_pie:
        run_pie = parrafo_pie.add_run(configuracion.texto_pie)
        run_pie.font.name = configuracion.fuente
        run_pie.font.size = Pt(8)
        run_pie.font.color.rgb = RGBColor(100, 100, 100)

    if configuracion.numeracion_paginas:
        if configuracion.texto_pie:
            parrafo_pie.add_run("   |   ")
        parrafo_pie.add_run("Página ")
        _agregar_campo_word(parrafo_pie, "PAGE")


def _agregar_portada(
    documento: Document,
    titulo: str,
    configuracion: ConfiguracionDocumento,
) -> None:
    """Crear una portada sencilla y editable cuando se solicite."""

    documento.add_paragraph("\n\n")

    titulo_portada = documento.add_paragraph()
    titulo_portada.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo_portada.paragraph_format.space_before = Pt(80)
    run_titulo = titulo_portada.add_run(titulo)
    run_titulo.bold = True
    run_titulo.font.name = configuracion.fuente
    run_titulo.font.size = Pt(28)
    run_titulo.font.color.rgb = RGBColor(31, 78, 121)

    if configuracion.subtitulo_portada:
        p_subtitulo = documento.add_paragraph()
        p_subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_subtitulo = p_subtitulo.add_run(configuracion.subtitulo_portada)
        run_subtitulo.font.name = configuracion.fuente
        run_subtitulo.font.size = Pt(16)

    if configuracion.institucion:
        p_institucion = documento.add_paragraph()
        p_institucion.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_institucion.paragraph_format.space_before = Pt(36)
        p_institucion.add_run(configuracion.institucion)

    if configuracion.autor:
        p_autor = documento.add_paragraph()
        p_autor.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_autor.paragraph_format.space_before = Pt(18)
        p_autor.add_run(configuracion.autor)

    documento.add_page_break()


# ==========================================
# GENERADOR PROFESIONAL DE CONTENIDO
# ==========================================

def _alineacion(nombre: Optional[str]) -> Optional[int]:
    """Convertir una alineación legible a la constante de Word."""

    if nombre is None:
        return None

    opciones = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "izquierda": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "centro": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "derecha": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "justificado": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }

    clave = nombre.strip().lower()
    if clave not in opciones:
        raise ValueError(f"Alineación no reconocida: {nombre}")
    return opciones[clave]


def _texto_bloque(bloque: Mapping[str, Any], clave: str = "texto") -> str:
    """Obtener y validar el texto de un bloque de contenido."""

    valor = bloque.get(clave, "")
    if not isinstance(valor, str):
        raise TypeError(f"El campo '{clave}' de un bloque debe ser texto.")
    return valor


def _agregar_bloque(documento: Document, bloque: Any, configuracion: ConfiguracionDocumento) -> None:
    """Añadir al documento un bloque estructurado."""

    if isinstance(bloque, str):
        if bloque.strip():
            _agregar_parrafo(documento, bloque)
        return

    tipo = str(bloque.get("tipo", bloque.get("type", "paragraph"))).strip().lower()

    if tipo in {"paragraph", "parrafo", "párrafo", "texto"}:
        partes = bloque.get("partes")
        if partes is not None:
            if not isinstance(partes, list):
                raise TypeError("El campo 'partes' de un párrafo debe ser una lista.")
            parrafo = documento.add_paragraph()
            for parte in partes:
                if not isinstance(parte, Mapping):
                    raise TypeError("Cada parte de un párrafo debe ser un diccionario.")
                p_tipo = str(parte.get("tipo", "texto")).strip().lower()
                if p_tipo == "texto":
                    parrafo.add_run(parte.get("contenido", ""))
                elif p_tipo == "ecuacion":
                    latex = parte.get("latex", "")
                    if not latex:
                        raise ValueError("Una parte de ecuación necesita el campo 'latex'.")
                    try:
                        omml = latex_a_omml(latex)
                    except LatexConversionError as error:
                        raise ValueError(
                            f"No se pudo convertir la ecuación LaTeX: {error}"
                        )
                    run = parrafo.add_run()
                    run._r.append(omml)
                else:
                    raise ValueError(f"Tipo de parte de párrafo no reconocido: {p_tipo}")
            alineacion = _alineacion(bloque.get("alineacion", bloque.get("align")))
            if alineacion is not None:
                parrafo.alignment = alineacion
            if bloque.get("negrita") or bloque.get("bold"):
                for run in parrafo.runs:
                    run.bold = True
            if bloque.get("cursiva") or bloque.get("italic"):
                for run in parrafo.runs:
                    run.italic = True
            return
        texto = _texto_bloque(bloque)
        if not texto.strip():
            return
        parrafo = documento.add_paragraph(texto)
        alineacion = _alineacion(bloque.get("alineacion", bloque.get("align")))
        if alineacion is not None:
            parrafo.alignment = alineacion
        if bloque.get("negrita") or bloque.get("bold"):
            for run in parrafo.runs:
                run.bold = True
        if bloque.get("cursiva") or bloque.get("italic"):
            for run in parrafo.runs:
                run.italic = True
        return

    if tipo in {"heading", "encabezado", "titulo", "título"}:
        texto = _texto_bloque(bloque)
        nivel = bloque.get("nivel", bloque.get("level", 1))
        if not isinstance(nivel, int) or nivel not in (1, 2, 3):
            raise ValueError("El nivel de encabezado debe ser 1, 2 o 3.")
        documento.add_heading(texto, level=nivel)
        return

    if tipo in {"bullet", "viñeta", "vineta", "lista_viñetas", "lista_vinetas"}:
        texto = _texto_bloque(bloque)
        documento.add_paragraph(texto, style="List Bullet")
        return

    if tipo in {"numbered", "numerado", "lista_numerada"}:
        texto = _texto_bloque(bloque)
        documento.add_paragraph(texto, style="List Number")
        return

    if tipo in {"quote", "cita"}:
        texto = _texto_bloque(bloque)
        documento.add_paragraph(texto, style="Quote")
        return

    if tipo in {"page_break", "salto_pagina", "salto_de_pagina"}:
        documento.add_page_break()
        return

    if tipo in {"section_break", "salto_seccion", "seccion", "section"}:
        orientacion = bloque.get("orientacion", bloque.get("orientation", "vertical"))
        if not isinstance(orientacion, str):
            raise TypeError("La orientación de una sección debe ser texto.")
        orientacion = orientacion.strip().lower()
        if orientacion not in {"vertical", "portrait", "horizontal", "landscape"}:
            raise ValueError("La orientación de sección debe ser vertical o horizontal.")
        nueva_seccion = documento.add_section(WD_SECTION.NEW_PAGE)
        _configurar_seccion(nueva_seccion, configuracion, orientacion)
        return

    if tipo in {"equation", "ecuacion"}:
        modo = str(bloque.get("modo", "display")).strip().lower()
        latex = bloque.get("latex", "")
        if not latex:
            raise ValueError("El bloque de ecuación necesita el campo 'latex'.")
        try:
            if modo == "inline":
                omml = latex_a_omml(latex)
                parrafo = documento.add_paragraph()
                # Eliminar el run vacío por defecto del párrafo nuevo.
                for child in list(parrafo._p):
                    parrafo._p.remove(child)
                run = parrafo.add_run()
                run._r.append(omml)
            else:
                omml_para = latex_a_omml_parrafo(latex)
                parrafo = documento.add_paragraph()
                for child in list(parrafo._p):
                    parrafo._p.remove(child)
                parrafo._p.append(omml_para)
                parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except LatexConversionError as error:
            raise ValueError(f"No se pudo convertir la ecuación LaTeX: {error}")
        return

    if tipo in {"image", "imagen"}:
        ruta_imagen = bloque.get("ruta", bloque.get("path", bloque.get("imagen")))
        if not ruta_imagen:
            raise ValueError("El bloque de imagen necesita el campo 'ruta' o 'path'.")
        archivo_imagen = Path(ruta_imagen)
        if not archivo_imagen.is_file():
            raise FileNotFoundError(f"No se encontró la imagen: {archivo_imagen}")

        ancho = bloque.get("ancho", bloque.get("width_inches", 5.5))
        if not isinstance(ancho, (int, float)) or ancho <= 0:
            raise ValueError("El ancho de la imagen debe ser mayor que cero.")

        parrafo_imagen = documento.add_paragraph()
        parrafo_imagen.alignment = _alineacion(
            bloque.get("alineacion", bloque.get("align", "center"))
        )
        parrafo_imagen.add_run().add_picture(str(archivo_imagen), width=Inches(float(ancho)))

        leyenda = bloque.get("leyenda", bloque.get("caption"))
        if leyenda:
            if not isinstance(leyenda, str):
                raise TypeError("La leyenda de una imagen debe ser texto.")
            parrafo_leyenda = documento.add_paragraph(leyenda, style="Caption")
            parrafo_leyenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return

    if tipo in {"table", "tabla"}:
        encabezados = bloque.get("encabezados", bloque.get("headers", []))
        filas = bloque.get("filas", bloque.get("rows", []))

        if not isinstance(encabezados, (list, tuple)):
            raise TypeError("Los encabezados de la tabla deben ser una lista.")
        if not isinstance(filas, (list, tuple)):
            raise TypeError("Las filas de la tabla deben ser una lista.")
        if not encabezados and not filas:
            raise ValueError("La tabla debe contener encabezados o filas.")

        cantidad_columnas = len(encabezados) if encabezados else len(filas[0])
        if cantidad_columnas < 1:
            raise ValueError("La tabla debe tener al menos una columna.")

        for indice, fila in enumerate(filas):
            if not isinstance(fila, (list, tuple)):
                raise TypeError(f"La fila {indice + 1} debe ser una lista.")
            if len(fila) != cantidad_columnas:
                raise ValueError(
                    f"La fila {indice + 1} tiene {len(fila)} celdas; "
                    f"se esperaban {cantidad_columnas}."
                )

        tabla = documento.add_table(
            rows=1 if encabezados else 0,
            cols=cantidad_columnas,
        )
        estilo_tabla = bloque.get("estilo", bloque.get("style", "Table Grid"))
        try:
            tabla.style = estilo_tabla
        except KeyError:
            tabla.style = "Table Grid"
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabla.autofit = True

        if encabezados:
            for indice, valor in enumerate(encabezados):
                celda = tabla.rows[0].cells[indice]
                celda.text = str(valor)
                celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for parrafo in celda.paragraphs:
                    for run in parrafo.runs:
                        run.bold = True

        for fila in filas:
            celdas = tabla.add_row().cells
            for indice, valor in enumerate(fila):
                celdas[indice].text = str(valor)
                celdas[indice].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        documento.add_paragraph()
        return

    raise ValueError(f"Tipo de bloque no reconocido: {tipo}")


def crear_documento_profesional(
    titulo: str,
    bloques: Iterable[Any],
    ruta_destino: Optional[RUTA_DESTINO] = None,
    configuracion: Optional[ConfiguracionDocumento] = None,
) -> Path:
    """Crear un documento DOCX profesional con contenido estructurado.

    Cada elemento de ``bloques`` puede ser texto o un diccionario. Tipos
    disponibles: ``paragraph``, ``heading``, ``bullet``, ``numbered``,
    ``quote``, ``table``, ``image``, ``equation``, ``page_break`` y
    ``section_break``.

    Ejemplo de ecuación independiente::

        {"tipo": "equation", "modo": "display", "latex": "\\int_0^1 x^2\\,dx"}

    Ejemplo de tabla::

        {"tipo": "tabla", "encabezados": ["Materia", "Nota"],
         "filas": [["Matemática", "18"], ["Programación", "20"]]}

    Ejemplo para cambiar a una página horizontal::

        {"tipo": "section_break", "orientacion": "horizontal"}

    El método mantiene el texto, las tablas, las imágenes y las ecuaciones
    como elementos editables de Word siempre que el formato lo permita. Las
    ecuaciones LaTeX se convierten a Office Math Markup Language (OMML), el
    formato nativo de ecuaciones de Microsoft Word.
    """

    titulo, elementos = _validar_bloques(titulo, bloques)
    opciones = configuracion or ConfiguracionDocumento()
    _validar_configuracion(opciones)

    documento = Document()
    _aplicar_estilos_profesionales(documento, opciones)
    _configurar_seccion(documento.sections[0], opciones)

    propiedades = documento.core_properties
    propiedades.title = titulo
    if opciones.autor:
        propiedades.author = opciones.autor
    if opciones.asunto:
        propiedades.subject = opciones.asunto
    if opciones.palabras_clave:
        propiedades.keywords = opciones.palabras_clave

    if opciones.crear_portada:
        _agregar_portada(documento, titulo, opciones)
    else:
        _agregar_titulo(documento, titulo)
        # El generador básico fija Arial por compatibilidad; la variante
        # profesional respeta la fuente seleccionada por el usuario.
        for run in documento.paragraphs[0].runs:
            run.font.name = opciones.fuente

    for bloque in elementos:
        _agregar_bloque(documento, bloque, opciones)

    _configurar_encabezado_pie(documento, opciones)

    ruta = _resolver_destino(titulo, ruta_destino)
    documento.save(ruta)
    return ruta


# ==========================================
# API PÚBLICA EXISTENTE
# ==========================================

def crear_documento(
    titulo: str,
    parrafos: Iterable[str],
    ruta_destino: Optional[RUTA_DESTINO] = None,
    *,
    procesar_marcas: bool = False,
) -> Path:
    """Crear un documento Word básico y devolver la ruta del archivo.

    ``procesar_marcas`` activa la interpretación de ``# ``, ``## ``,
    ``### ``, ``- ``, ``* `` y ``N. ``/``N) ``. Sin él, todo el contenido
    se escribe como párrafos normales.

    Nunca sobrescribe: si el destino existe, añade ``_2``, ``_3``…
    """

    titulo, textos = _validar_entrada(titulo, parrafos)

    documento = Document()
    seccion = _aplicar_formato(documento)

    _agregar_titulo(documento, titulo)
    _agregar_contenido(documento, textos, procesar_marcas)

    seccion.footer.paragraphs[0].text = "Generado con Word-Architect"

    ruta = _resolver_destino(titulo, ruta_destino)
    documento.save(ruta)

    return ruta
